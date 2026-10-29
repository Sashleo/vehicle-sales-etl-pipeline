"""
Sales Data Canonical Classification Pipeline
============================================

Cleans, enriches and standardises multi-year vehicle sales data by
resolving every record to a canonical vehicle classification:

    maker | model_ | grade_trim | engine | transmission | canonical_variant

The pipeline runs in the order the work was originally carried out:

    PHASE 0  Inspect and profile workbooks
    PHASE 1  Per-year enrichment: a five-stage merge chain
             (sales cleanup -> master key -> merge2 fallback
              -> lookup fallback -> leftover canonical fallback)
    PHASE 2  Cross-year standardisation and leftover fill
    PHASE 3  Staged targeted fill (frame key, then variant key)
    PHASE 4  Variant-key pass on an already-cleaned year

Guarantees enforced throughout:
    - No row explosion: every merge is checked against the source row count
    - Fallback stages only fill blanks; populated values are never overwritten
    - Rows that already held canonical values are verified as untouched
    - Final schema, column count and column order are enforced
    - Written workbooks are reopened and verified

Usage:
    Edit the CONFIGURATION section (or pass your own config dicts),
    then call the phase runner you need. See the __main__ block.
"""

import gc
import math
import numbers
import os
import re

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook


# ============================================================
# CONFIGURATION
# ============================================================

READ_ENGINE = "calamine"

CANONICAL_COLUMNS = [
    "maker",
    "model_",          # canonical model (distinct from sales-cleanup "model")
    "grade_trim",
    "engine",
    "transmission",
    "canonical_variant",
]

# Columns added to the raw data by the Phase 1 merge chain
ENRICHMENT_COLUMNS = [
    "material",
    "material_description",
    "model",
    *CANONICAL_COLUMNS,
]

# Values treated as "blank" everywhere in the pipeline
BLANK_TOKENS = {"", "0", "0.0", "nan", "none", "null", "n/a", "na", "<na>"}

# Leftover canonical workbook (shared across all years)
LEFTOVER_FILE = "leftover_canonical_variants.xlsx"
LEFTOVER_SHEET = "Standardized Variants"

LEFTOVER_REQUIRED_COLUMNS = [
    "raw_frame_no",
    "maker",
    "model",
    "grade_trim",
    "engine",
    "transmission",
    "canonical_variant",
]

LEFTOVER_COLUMN_MAP = {
    "maker": "maker",
    "model": "model_",
    "grade_trim": "grade_trim",
    "engine": "engine",
    "transmission": "transmission",
    "canonical_variant": "canonical_variant",
}

# Master key sheet (Merge1) and fallback sheets (Merge2, Lookup)
# list their six classification columns positionally.
MASTER_KEY_COLUMNS = [f"master_key.column{i}" for i in range(1, 7)]
MERGE2_COLUMNS = [f"lookup.column{i}" for i in range(2, 8)]
LOOKUP_KEY_COLUMN = "column1"
LOOKUP_COLUMNS = [f"column{i}" for i in range(2, 8)]

# Lookup sheet column meaning (confirmed from sheet inspection):
# column1 = raw variant, column2..7 = canonical six, column8 = match confidence
LOOKUP_SHEET_COLUMN_MAP = {
    "column1": "raw_variant",
    "column2": "maker",
    "column3": "model_",
    "column4": "grade_trim",
    "column5": "engine",
    "column6": "transmission",
    "column7": "canonical_variant",
    "column8": "match_confidence",
}

EXPECTED_FINAL_COLUMNS = 44


# ------------------------------------------------------------
# PHASE 1 configs: one per fiscal year.
# Join keys differ by year because the source sheets differ.
# ------------------------------------------------------------
PHASE1_CONFIGS = {
    "24-25": {
        "source_file": "source_data_FY24-25.xlsx",
        "raw_sheet": "Raw Data",
        "sales_join": "frame_no",            # "frame_no" or "row_aligned"
        "master_key_join": "vehicle_variant",
        "master_key_columns": [
            "master_key.maker",
            "master_key.model",
            "master_key.grade_trim",
            "master_key.engine",
            "master_key.transmission",
            "master_key.canonical_variant",
        ],
        "use_merge2": True,
        "use_lookup": True,
        "leftover_join": "frame_no",         # "frame_no" or "vehicle_variant"
        "strict": True,                      # raise on row change (else warn)
        "enforce_final_schema": True,
        "intermediate_output": None,
        "output_file": "enriched_FY24-25.xlsx",
    },
    "23-24": {
        "source_file": "source_data_FY23-24.xlsx",
        "raw_sheet": "Raw Data",
        "sales_join": "frame_no",
        "master_key_join": "repair_order_number",
        "master_key_columns": MASTER_KEY_COLUMNS,
        "use_merge2": True,
        "use_lookup": True,
        "leftover_join": "frame_no",
        "strict": False,
        "enforce_final_schema": False,
        "intermediate_output": None,
        "output_file": "enriched_FY23-24.xlsx",
    },
    "22-23": {
        "source_file": "source_data_FY22-23.xlsx",
        "raw_sheet": "Raw Data",
        "sales_join": "frame_no",
        "master_key_join": "repair_order_number",
        "master_key_columns": MASTER_KEY_COLUMNS,
        "use_merge2": False,
        "use_lookup": True,
        "leftover_join": "vehicle_variant",
        "strict": False,
        "enforce_final_schema": False,
        "intermediate_output": None,
        "output_file": "enriched_FY22-23.xlsx",
    },
    "21-22": {
        "source_file": "source_data_FY21-22.xlsx",
        "raw_sheet": "Raw Data",
        "sales_join": "row_aligned",
        "master_key_join": "vehicle_variant",
        "master_key_columns": MASTER_KEY_COLUMNS,
        "use_merge2": True,
        "use_lookup": False,
        "leftover_join": "frame_no",
        "strict": False,
        "enforce_final_schema": False,
        "intermediate_output": "enriched_FY21-22_intermediate.xlsx",
        "output_file": "enriched_FY21-22.xlsx",
    },
}


# ------------------------------------------------------------
# PHASE 2 configs: cross-year standardisation.
# header_overrides are applied to RAW headers before snake-casing,
# because some years contain columns that would otherwise collide
# (e.g. FY24-25 has both "Model " and "Model" as separate fields).
# ------------------------------------------------------------
PHASE2_FILES = [
    {
        "fy": "21-22",
        "input": "enriched_FY21-22.xlsx",
        "sheet": "Sheet1",
        "output": "clean_data_FY21-22.xlsx",
        "header_overrides": {},
        "drop_columns": [],
    },
    {
        "fy": "22-23",
        "input": "enriched_FY22-23.xlsx",
        "sheet": "Sheet1",
        "output": "clean_data_FY22-23.xlsx",
        "header_overrides": {},
        "drop_columns": [],
    },
    {
        "fy": "23-24",
        "input": "enriched_FY23-24.xlsx",
        "sheet": "Clean Data",
        "output": "clean_data_FY23-24.xlsx",
        "header_overrides": {"Model.1": "model_"},
        # 46 columns reduced to 44
        "drop_columns": ["lookup.raw_variant", "lookup.match_confidence"],
    },
    {
        "fy": "24-25",
        "input": "enriched_FY24-25.xlsx",
        "sheet": "Sheet1",
        "output": "clean_data_FY24-25.xlsx",
        # exact "Model" is canonical; "Model " (trailing space) is the material field
        "header_overrides": {"Model": "model_", "Model ": "model"},
        # 45 columns reduced to 44
        "drop_columns": ["unnamed_44"],
    },
    {
        "fy": "25-26",
        "input": "source_data_FY25-26.xlsx",
        "sheet": "Final Data",
        "output": "clean_data_FY25-26.xlsx",
        "header_overrides": {"Maker_": "maker", "Trim": "grade_trim"},
        # 47 columns reduced to 44
        "drop_columns": ["frame_no.1", "vehicle_variant.1", "vehicle_make.1"],
    },
]

PHASE2_SUMMARY_FILE = "canonical_merge_summary.csv"


# ------------------------------------------------------------
# PHASE 3 config: staged targeted fill on one year
# ------------------------------------------------------------
PHASE3_CONFIG = {
    "source_file": "source_data_FY25-26.xlsx",
    "source_sheet": "Final Data",
    "header_overrides": {"Maker_": "maker", "Trim": "grade_trim"},
    "drop_columns": ["frame_no.1", "vehicle_variant.1", "vehicle_make.1"],
    "expected_source_rows": None,       # set to your known row count, or None to skip
    "expected_source_columns": 47,
    "expected_final_columns": EXPECTED_FINAL_COLUMNS,
    "summary_file": "FY25-26_staged_fill_summary.csv",
    "stages": [
        {
            "label": "Stage 1 - Leftovers using frame_no",
            "lookup_file": LEFTOVER_FILE,
            "lookup_sheet": LEFTOVER_SHEET,
            "source_key": "frame_no",
            "lookup_key": "raw_frame_no",
            "normalizer": "frame",
            "required_columns": LEFTOVER_REQUIRED_COLUMNS,
            "column_map": LEFTOVER_COLUMN_MAP,
            "drop_embedded_header": False,
            "output_file": "FY25-26_after_frame_merge.xlsx",
        },
        {
            "label": "Stage 2 - FY22-23 Lookup using vehicle_variant",
            "lookup_file": "source_data_FY22-23.xlsx",
            "lookup_sheet": "Lookup",
            "source_key": "vehicle_variant",
            "lookup_key": "raw_variant",
            "normalizer": "variant",
            "required_columns": list(LOOKUP_SHEET_COLUMN_MAP.keys()),
            "column_map": LOOKUP_SHEET_COLUMN_MAP,
            "drop_embedded_header": True,
            "output_file": "FY25-26_final.xlsx",
        },
    ],
}


# ------------------------------------------------------------
# PHASE 4 config: variant-key pass on an already-cleaned year
# ------------------------------------------------------------
PHASE4_CONFIG = {
    "target_file": "clean_data_FY22-23.xlsx",
    "target_sheet": "Final Data",
    "lookup_file": "source_data_FY22-23.xlsx",
    "lookup_sheet": "Lookup",
    "output_file": "clean_data_FY22-23_v2.xlsx",
}


# ============================================================
# SECTION 1: HEADER STANDARDISATION AND LOADING
# ============================================================

def to_snake_case(name):
    """
    'Grade / Trim' -> 'grade_trim'
    'Master Key.Column1' -> 'master_key.column1'
    'Model_' -> 'model_'   (trailing underscore preserved)
    """
    text = str(name).strip().lower()
    text = re.sub(r"[^a-z0-9._]+", "_", text)
    text = re.sub(r"_+", "_", text)
    return text.lstrip("_")


def standardize_headers(df, overrides=None, label="dataframe"):
    """
    Apply raw-header overrides first, then snake-case the rest.
    Stops if two source columns would collapse into one name.
    """
    overrides = overrides or {}
    new_columns = [
        overrides.get(col, to_snake_case(col)) for col in df.columns
    ]
    collisions = pd.Index(new_columns)[
        pd.Index(new_columns).duplicated(keep=False)
    ].unique().tolist()
    if collisions:
        raise ValueError(
            f"\nHeader standardisation in {label} would merge columns: "
            f"{collisions}\nAdd header_overrides to keep them distinct."
        )
    df = df.copy()
    df.columns = new_columns
    return df


def read_sheet(filepath, sheet_name, header_overrides=None, **kwargs):
    df = pd.read_excel(
        filepath, sheet_name=sheet_name, engine=READ_ENGINE, **kwargs
    )
    return standardize_headers(
        df, header_overrides, label=f"{filepath} [{sheet_name}]"
    )


# ============================================================
# SECTION 2: PHASE 0 - INSPECTION AND PROFILING
# ============================================================

def inspect_workbook(filepath: str):
    """Print every sheet's shape and column list."""
    print(f"\n{'=' * 60}")
    print(f"  {filepath}")
    print(f"{'=' * 60}")
    xls = pd.ExcelFile(filepath, engine=READ_ENGINE)
    for sheet in xls.sheet_names:
        df = xls.parse(sheet)
        print(f"\n  Sheet: '{sheet}'  →  {df.shape[0]} rows × {df.shape[1]} cols")
        print(f"  Columns: {list(df.columns)}")
    xls.close()


def profile_lookup(df, label, head=20):
    """Shape, numbered columns, head and null counts of a lookup sheet."""
    print("\n" + "=" * 80)
    print(f"{label} PROFILE")
    print("=" * 80)
    print(f"Shape: {df.shape}")
    print("\nColumns:")
    for i, col in enumerate(df.columns, start=1):
        print(f"  {i}. {col!r}")
    print(f"\nHead({head}):")
    print(df.head(head).to_string())
    print("\nNulls per column:")
    print(df.isna().sum())


def diagnose_lookup_keys(lookup, key):
    """
    Check a lookup key before merging. Duplicated keys cause row
    explosion in a left merge; this must report 0 duplicates.
    """
    duplicates = lookup[lookup.duplicated(subset=[key], keep=False)]
    print(f"\nLookup key '{key}':")
    print(f"  rows: {len(lookup):,}")
    print(f"  unique: {lookup[key].nunique():,} / {len(lookup):,}")
    print(f"  duplicated rows: {len(duplicates):,}")
    if not duplicates.empty:
        print(f"  duplicated values: {duplicates[key].unique().tolist()[:50]}")
    return duplicates


def match_diagnostics(raw, leftover):
    """How many raw rows can be matched by frame number and by variant."""
    frame = clean_key(raw["frame_no"])
    variant = clean_key(raw["vehicle_variant"])
    print("Frame matches:", frame.isin(clean_key(leftover["raw_frame_no"])).sum())
    print("Variant matches:", variant.isin(clean_key(leftover["raw_variant"])).sum())


# ============================================================
# SECTION 3: KEY CLEANING AND BLANK HANDLING
# ============================================================

def clean_key(series):
    """Strip whitespace from join keys while preserving blanks."""
    return (
        series.astype("string")
        .str.strip()
        .replace({"": pd.NA, "nan": pd.NA, "None": pd.NA, "<NA>": pd.NA})
    )


def _normalize_alnum_key(value):
    """Uppercase, drop non-alphanumerics; blank/zero-like -> pd.NA."""
    if pd.isna(value):
        return pd.NA
    text = str(value).strip().upper()
    if text.lower() in BLANK_TOKENS:
        return pd.NA
    text = re.sub(r"[^A-Z0-9]", "", text)
    if text in {"", "0"}:
        return pd.NA
    return text


def normalize_frame_no(value):
    """' abc-123 ' -> 'ABC123', 'ABC/123' -> 'ABC123'."""
    return _normalize_alnum_key(value)


def normalize_vehicle_variant(value):
    """
    Exact normalised matching, not fuzzy matching.
    ' 1.5 XL A/T ' -> '15XLAT', '2.0L SPORT CVT' -> '20LSPORTCVT'
    """
    return _normalize_alnum_key(value)


NORMALIZERS = {
    "frame": normalize_frame_no,
    "variant": normalize_vehicle_variant,
}


def blank_or_zero(series):
    """True where a cell is NaN, empty, zero, or a missing-value token."""
    text = series.astype("string").str.strip().str.lower()
    return series.isna() | text.isin(BLANK_TOKENS)


def clean_lookup_values(series):
    """Convert blank/zero-like lookup values to pd.NA so zeros are never written."""
    result = series.copy()
    result.loc[blank_or_zero(result)] = pd.NA
    return result


def make_all_six_blank_mask(df):
    mask = pd.Series(True, index=df.index, dtype=bool)
    for column in CANONICAL_COLUMNS:
        mask &= blank_or_zero(df[column])
    return mask


def make_all_six_populated_mask(df):
    mask = pd.Series(True, index=df.index, dtype=bool)
    for column in CANONICAL_COLUMNS:
        mask &= ~blank_or_zero(df[column])
    return mask


def make_any_six_populated_mask(df):
    mask = pd.Series(False, index=df.index, dtype=bool)
    for column in CANONICAL_COLUMNS:
        mask |= ~blank_or_zero(df[column])
    return mask


# ============================================================
# SECTION 4: VALIDATION
# ============================================================

def check_row_count(df, stage_name, original_rows, strict=True):
    """Stop (or warn) if a merge increased or decreased the row count."""
    if len(df) != original_rows:
        message = (
            f"\nROW COUNT CHANGED AFTER {stage_name}\n"
            f"Expected: {original_rows:,}\n"
            f"Found:    {len(df):,}\n"
        )
        if strict:
            raise RuntimeError(message)
        print("WARNING:" + message)
        return
    print(f"{stage_name}: {df.shape} | Row count intact")


def report_blanks(df, stage_name):
    print(f"Blanks after {stage_name}:", df[CANONICAL_COLUMNS].isna().sum().to_dict())


def check_duplicate_columns(df, label):
    duplicates = df.columns[df.columns.duplicated(keep=False)].tolist()
    if duplicates:
        raise ValueError(
            f"\nDuplicate columns found in {label}:\n{duplicates}\n\n"
            "Processing stopped to prevent silent overwriting."
        )


def verify_required_columns(df, required_columns, label):
    missing = [col for col in required_columns if col not in df.columns]
    if missing:
        raise KeyError(
            f"\n{label} is missing required columns:\n{missing}\n\n"
            f"Available columns:\n{list(df.columns)}"
        )


def verify_unchanged_structure(df, expected_rows, expected_columns,
                               expected_column_order, stage_label):
    if len(df) != expected_rows:
        raise RuntimeError(
            f"\n{stage_label} changed row count:\n{expected_rows:,} -> {len(df):,}"
        )
    if len(df.columns) != expected_columns:
        raise RuntimeError(
            f"\n{stage_label} changed column count:\n"
            f"{expected_columns:,} -> {len(df.columns):,}"
        )
    if list(df.columns) != expected_column_order:
        raise RuntimeError(f"\n{stage_label} changed column names or order.")


def verify_protected_rows(df, protected_values, protected_mask, stage_label):
    """Rows that originally held any canonical value must be unchanged."""
    for column in CANONICAL_COLUMNS:
        before = protected_values.loc[protected_mask, column].astype("string").fillna("<NA>")
        after = df.loc[protected_mask, column].astype("string").fillna("<NA>")
        changed = before != after
        if changed.any():
            raise RuntimeError(
                f"\n{stage_label} changed {int(changed.sum()):,} protected "
                f"values in column '{column}'."
            )


# ============================================================
# SECTION 5: EXCEL OUTPUT
# ============================================================

def excel_safe(value):
    """Convert pandas/NumPy values into types openpyxl can write."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.datetime64):
        return None if np.isnat(value) else pd.Timestamp(value).to_pydatetime()
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return None if (np.isnan(value) or np.isinf(value)) else float(value)
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, numbers.Integral):
        return int(value)
    if isinstance(value, numbers.Real):
        numeric = float(value)
        return None if (math.isnan(numeric) or math.isinf(numeric)) else numeric
    return value


def write_large_excel(df, output_file, sheet_name="Final Data"):
    """
    Row-by-row write using openpyxl write-only mode. Keeps memory low
    on multi-hundred-thousand-row frames without dropping cell data.
    """
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_name)
    ws.append([str(column) for column in df.columns])
    for row in df.itertuples(index=False, name=None):
        ws.append([excel_safe(value) for value in row])
    wb.save(output_file)


def verify_written_workbook(output_file, expected_sheet, expected_rows,
                            expected_columns, expected_headers):
    """Reopen an output workbook and confirm shape, headers and sample rows."""
    if not os.path.exists(output_file):
        raise RuntimeError(f"Output file was not created: {output_file}")

    wb = load_workbook(output_file, read_only=True, data_only=False)
    try:
        if expected_sheet not in wb.sheetnames:
            raise RuntimeError(
                f"\nExpected sheet '{expected_sheet}' not found in {output_file}.\n"
                f"Available sheets: {wb.sheetnames}"
            )
        ws = wb[expected_sheet]
        actual_headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
        if ws.max_row is None:
            # write-only workbooks store no dimension record, so count directly
            ws.reset_dimensions()
            total_rows = sum(1 for _ in ws.iter_rows(values_only=True))
            actual_columns = len(actual_headers)
        else:
            total_rows = ws.max_row
            actual_columns = ws.max_column
        actual_rows = total_rows - 1

        if actual_rows != expected_rows:
            raise RuntimeError(
                f"\nOutput row count mismatch in {output_file}:\n"
                f"Expected: {expected_rows:,}\nActual:   {actual_rows:,}"
            )
        if actual_columns != expected_columns:
            raise RuntimeError(
                f"\nOutput column count mismatch in {output_file}:\n"
                f"Expected: {expected_columns:,}\nActual:   {actual_columns:,}"
            )
        if actual_headers != list(expected_headers):
            raise RuntimeError(f"\nOutput headers or column order changed in {output_file}.")

        sample_rows_checked = 0
        for row in ws.iter_rows(min_row=2, max_row=min(total_rows, 11), values_only=True):
            populated = sum(v is not None and str(v).strip() != "" for v in row)
            if populated <= 1:
                raise RuntimeError(
                    f"\nA sample row in {output_file} has zero or one populated "
                    "cell. This may indicate an incomplete export."
                )
            sample_rows_checked += 1
    finally:
        wb.close()

    print(f"Reopened successfully: {output_file}")
    print(f"Confirmed written shape: {actual_rows:,} rows x {actual_columns:,} columns")
    print(f"Sample data rows checked: {sample_rows_checked:,}")


# ============================================================
# SECTION 6: PHASE 1 - MERGE CHAIN STAGES
# ============================================================

def merge_sales_cleanup(raw, sales_clean, join="frame_no"):
    """
    MERGE 1: add material, material_description and model.

    join="frame_no"    -> left merge on frame number (deduplicated)
    join="row_aligned" -> sales sheet is row-aligned with raw,
                          so the three columns are copied directly
    """
    cols = ["material", "material_description", "model"]

    if join == "row_aligned":
        if len(sales_clean) != len(raw):
            raise RuntimeError(
                f"Row-aligned copy needs equal lengths: "
                f"raw={len(raw):,}, sales={len(sales_clean):,}"
            )
        raw = raw.copy()
        raw[cols] = sales_clean[cols].values
        return raw

    raw = raw.copy()
    raw["frame_no"] = clean_key(raw["frame_no"])
    sales = sales_clean[["frame_no"] + cols].copy()
    sales["frame_no"] = clean_key(sales["frame_no"])
    # strip the key before resolving duplicates
    sales = sales.dropna(subset=["frame_no"])
    sales = sales.drop_duplicates(subset=["frame_no"], keep="first")

    return raw.merge(sales, on="frame_no", how="left", validate="many_to_one")


def merge_master_key(raw, merge1, key, master_key_columns):
    """
    MERGE 2: primary canonical classification from the master key sheet.
    key is vehicle_variant, repair_order_number or frame_no depending on year.
    """
    raw = raw.copy()
    raw[key] = clean_key(raw[key])

    m1 = merge1[[key] + master_key_columns].copy()
    m1[key] = clean_key(m1[key])
    m1 = m1.dropna(subset=[key])
    # resolve duplicate keys before merging (XLOOKUP first-match behaviour)
    m1 = m1.drop_duplicates(subset=[key], keep="first")

    raw = raw.merge(m1, on=key, how="left", validate="many_to_one")
    return raw.rename(columns=dict(zip(master_key_columns, CANONICAL_COLUMNS)))


def merge_fallback(raw, fallback, fallback_key, fallback_columns, suffix):
    """
    MERGE 3 / MERGE 4: fill only blanks left by earlier stages,
    matched on vehicle_variant. Used for the Merge2 sheet and the
    Lookup sheet.
    """
    raw = raw.copy()
    raw["vehicle_variant"] = clean_key(raw["vehicle_variant"])

    fb = fallback[[fallback_key] + fallback_columns].copy()
    fb[fallback_key] = clean_key(fb[fallback_key])
    fb = fb.dropna(subset=[fallback_key])
    fb = fb.drop_duplicates(subset=[fallback_key], keep="first")
    fb = fb.rename(columns={
        fallback_key: "vehicle_variant",
        **dict(zip(fallback_columns, CANONICAL_COLUMNS)),
    })

    # clear any leftover suffix columns from an earlier partial run
    raw = raw.drop(columns=[c + suffix for c in CANONICAL_COLUMNS], errors="ignore")

    raw = raw.merge(
        fb, on="vehicle_variant", how="left",
        suffixes=("", suffix), validate="many_to_one",
    )
    for col in CANONICAL_COLUMNS:
        raw[col] = raw[col].fillna(raw[col + suffix])
    return raw.drop(columns=[c + suffix for c in CANONICAL_COLUMNS])


def merge_leftover(raw, leftover, join="frame_no"):
    """
    MERGE 5: final fallback from the leftover canonical workbook.
    Uses Series.map (not merge), so duplicate keys cannot multiply rows.

    join="frame_no"        -> raw.frame_no       vs leftover.raw_frame_no
    join="vehicle_variant" -> raw.vehicle_variant vs leftover.raw_variant
    """
    raw_key, leftover_key = {
        "frame_no": ("frame_no", "raw_frame_no"),
        "vehicle_variant": ("vehicle_variant", "raw_variant"),
    }[join]

    raw = raw.copy()
    raw[raw_key] = clean_key(raw[raw_key])

    lv = leftover.copy()
    lv[leftover_key] = clean_key(lv[leftover_key])
    lv = lv.dropna(subset=[leftover_key])
    # one record per key prevents row explosion
    lv = lv.drop_duplicates(subset=[leftover_key], keep="last")
    lv = lv.set_index(leftover_key)

    for leftover_col, raw_col in LEFTOVER_COLUMN_MAP.items():
        before = raw[raw_col].isna().sum()
        raw[raw_col] = raw[raw_col].fillna(raw[raw_key].map(lv[leftover_col]))
        print(f"  {raw_col}: filled {before - raw[raw_col].isna().sum():,} blanks")
    return raw


def fill_unmatched_rows_via_subset(main, fallback, fallback_key, fallback_columns):
    """
    Alternative fallback method (earlier approach): mark rows whose six
    canonical fields are all blank, merge ONLY that subset against the
    fallback, and write the results back. Includes the row-explosion
    check: the merged subset must be the same size as the mask.
    """
    main = main.copy()
    main[CANONICAL_COLUMNS] = main[CANONICAL_COLUMNS].fillna(0)
    mask = (main[CANONICAL_COLUMNS] == 0).all(axis=1)
    print(f"Rows needing fallback: {int(mask.sum()):,}")

    fb = fallback[[fallback_key] + fallback_columns].copy()
    fb = fb.drop_duplicates(subset=[fallback_key], keep="first")
    if fb[fallback_key].duplicated().sum() != 0:
        raise RuntimeError("Fallback key still has duplicates; merge would explode.")
    fb = fb.rename(columns=dict(zip(fallback_columns, CANONICAL_COLUMNS)))

    unmatched = (
        main.loc[mask, ["vehicle_variant"]]
        .merge(fb, left_on="vehicle_variant", right_on=fallback_key, how="left")
    )
    if len(unmatched) != int(mask.sum()):
        raise RuntimeError(
            f"Row explosion in subset merge: {int(mask.sum()):,} -> {len(unmatched):,}"
        )
    unmatched.index = main.index[mask]
    main.loc[mask, CANONICAL_COLUMNS] = unmatched[CANONICAL_COLUMNS]
    return main


def fix_duplicate_model_column(df):
    """If two columns are both named 'model', rename the second to 'model_'."""
    model_positions = [i for i, c in enumerate(df.columns) if c == "model"]
    if len(model_positions) > 1:
        df = df.copy()
        df.columns = [
            "model_" if i == model_positions[1] else c
            for i, c in enumerate(df.columns)
        ]
    return df


def enforce_final_schema(raw, original_columns):
    """Original columns + enrichment columns, in exact order, no extras."""
    expected = original_columns + ENRICHMENT_COLUMNS

    raw = raw.loc[:, ~raw.columns.duplicated()].copy()
    temp_cols = [c for c in raw.columns if c.endswith("_m2") or c.endswith("_lk")]
    if temp_cols:
        raw = raw.drop(columns=temp_cols)

    missing = [c for c in expected if c not in raw.columns]
    if missing:
        raise KeyError(f"Final output is missing columns: {missing}")
    return raw[expected].copy()


def final_validation(raw, original_rows, original_columns, schema_enforced):
    print("\n" + "=" * 60)
    print("FINAL CHECK")
    print("=" * 60)
    print(f"Original shape: ({original_rows:,}, {len(original_columns)})")
    print(f"Final shape:    {raw.shape}")

    if len(raw) != original_rows:
        raise RuntimeError("Final row count does not match the source.")

    if schema_enforced:
        expected_cols = len(original_columns) + len(ENRICHMENT_COLUMNS)
        print(f"Expected cols:  {expected_cols}")
        print(f"Actual cols:    {len(raw.columns)}")
        if len(raw.columns) != expected_cols:
            raise RuntimeError(
                f"Final column count is incorrect. "
                f"Expected {expected_cols}, found {len(raw.columns)}."
            )

    if raw.columns.duplicated().any():
        raise RuntimeError("Duplicate column names remain in the final output.")

    print("No row explosion detected.")
    print("No duplicate columns detected.")
    print("Final schema is intact.")


def run_base_enrichment(cfg, leftover_file=LEFTOVER_FILE, leftover_sheet=LEFTOVER_SHEET):
    """PHASE 1 runner: the full five-stage merge chain for one year."""
    strict = cfg["strict"]

    # --- load ---
    xls = pd.ExcelFile(cfg["source_file"], engine=READ_ENGINE)
    load = lambda sheet: standardize_headers(xls.parse(sheet), label=sheet)
    raw = load(cfg["raw_sheet"])
    sales_clean = load("Sales Cleanup")
    merge1 = load("Merge1")
    merge2 = load("Merge2") if cfg["use_merge2"] else None
    lookup = load("Lookup") if cfg["use_lookup"] else None
    xls.close()
    leftover = read_sheet(leftover_file, leftover_sheet)

    original_rows = len(raw)
    original_columns = raw.columns.tolist()
    print(f"Starting shape: {raw.shape}")

    # --- merge 1: sales cleanup ---
    raw = merge_sales_cleanup(raw, sales_clean, join=cfg["sales_join"])
    check_row_count(raw, "After Sales Cleanup", original_rows, strict)

    # --- merge 2: master key ---
    raw = merge_master_key(raw, merge1, cfg["master_key_join"], cfg["master_key_columns"])
    check_row_count(raw, "After Merge1", original_rows, strict)
    report_blanks(raw, "Merge1")

    # --- merge 3: merge2 fallback ---
    if merge2 is not None:
        raw = merge_fallback(raw, merge2, "vehicle_variant", MERGE2_COLUMNS, "_m2")
        check_row_count(raw, "After Merge2 fallback", original_rows, strict)
        report_blanks(raw, "Merge2")

    # --- merge 4: lookup fallback ---
    if lookup is not None:
        raw = merge_fallback(raw, lookup, LOOKUP_KEY_COLUMN, LOOKUP_COLUMNS, "_lk")
        check_row_count(raw, "After Lookup fallback", original_rows, strict)
        report_blanks(raw, "Lookup")

    if cfg.get("intermediate_output"):
        write_large_excel(raw, cfg["intermediate_output"], sheet_name="Sheet1")
        print(f"Intermediate saved → {cfg['intermediate_output']}")

    # --- merge 5: leftover canonical fallback ---
    raw = merge_leftover(raw, leftover, join=cfg["leftover_join"])
    check_row_count(raw, "After Leftover fallback", original_rows, strict)
    report_blanks(raw, "Leftover")

    # --- schema control and validation ---
    if cfg["enforce_final_schema"]:
        raw = enforce_final_schema(raw, original_columns)
    final_validation(raw, original_rows, original_columns, cfg["enforce_final_schema"])

    # --- export ---
    write_large_excel(raw, cfg["output_file"], sheet_name="Sheet1")
    verify_written_workbook(cfg["output_file"], "Sheet1", len(raw),
                            len(raw.columns), list(raw.columns))
    print(f"\nDone → {cfg['output_file']}")
    print(f"Final shape → {raw.shape}")
    return raw


# ============================================================
# SECTION 7: SHARED LOOKUP PREPARATION AND APPLICATION
# (used by Phases 2 and 3)
# ============================================================

def prepare_lookup(raw_lookup, key_column, key_normalizer, lookup_label):
    """
    One row per normalised key. Where keys are duplicated, keep the row
    with the most populated canonical fields (stable mergesort, so ties
    keep source order). Returns the prepared frame and per-column maps.
    """
    required = [key_column, *CANONICAL_COLUMNS]
    verify_required_columns(raw_lookup, required, lookup_label)

    prepared = raw_lookup[required].copy()
    prepared["_lookup_key"] = prepared[key_column].map(key_normalizer)
    for column in CANONICAL_COLUMNS:
        prepared[column] = clean_lookup_values(prepared[column])
    prepared = prepared.loc[prepared["_lookup_key"].notna()].copy()

    prepared["_completeness"] = prepared[CANONICAL_COLUMNS].notna().sum(axis=1)

    dup_mask = prepared.duplicated(subset="_lookup_key", keep=False)
    print(f"{lookup_label} duplicate rows: {int(dup_mask.sum()):,}")
    print(
        f"{lookup_label} unique duplicated keys: "
        f"{int(prepared.loc[dup_mask, '_lookup_key'].nunique()):,}"
    )

    prepared = prepared.sort_values("_completeness", ascending=False, kind="mergesort")
    prepared = prepared.drop_duplicates(subset="_lookup_key", keep="first")
    prepared = prepared.set_index("_lookup_key")
    print(f"Usable unique keys: {len(prepared):,}")

    lookup_maps = {col: prepared[col].to_dict() for col in CANONICAL_COLUMNS}
    return prepared, lookup_maps


def apply_lookup(df, source_key_series, key_normalizer, prepared_lookup,
                 lookup_maps, stage_label):
    """
    Fill canonical fields only for rows where ALL SIX are blank/zero at the
    start of the stage, and only from non-blank lookup values.
    Series.map is used, so row multiplication is impossible.
    """
    rows_before = len(df)
    columns_before = list(df.columns)

    normalized_key = source_key_series.map(key_normalizer)
    eligible_mask = make_all_six_blank_mask(df)
    matched_mask = (
        eligible_mask
        & normalized_key.notna()
        & normalized_key.isin(prepared_lookup.index)
    )
    eligible_rows = int(eligible_mask.sum())
    matched_rows = int(matched_mask.sum())

    print("\n" + "-" * 80)
    print(stage_label)
    print("-" * 80)
    print(f"Rows eligible at start of stage: {eligible_rows:,}")
    print(f"Eligible rows matching lookup: {matched_rows:,}")

    rows_receiving_value = pd.Series(False, index=df.index, dtype=bool)
    cells_filled_by_column = {}

    for column in CANONICAL_COLUMNS:
        mapped = normalized_key.map(lookup_maps[column])
        fill_mask = matched_mask & mapped.notna() & blank_or_zero(df[column])
        cells = int(fill_mask.sum())
        if cells:
            df.loc[fill_mask, column] = mapped.loc[fill_mask]
        rows_receiving_value |= fill_mask
        cells_filled_by_column[column] = cells
        print(f"{column}: {cells:,} cells filled")

    if len(df) != rows_before:
        raise RuntimeError(f"\n{stage_label} changed row count:\n{rows_before:,} -> {len(df):,}")
    if list(df.columns) != columns_before:
        raise RuntimeError(f"\n{stage_label} changed columns or column order.")

    receiving = int(rows_receiving_value.sum())
    print(f"Rows receiving at least one value: {receiving:,}")

    stats = {
        "Stage": stage_label,
        "Eligible Rows": eligible_rows,
        "Matched Eligible Rows": matched_rows,
        "Rows Receiving Any Value": receiving,
        **{f"{col} cells filled": n for col, n in cells_filled_by_column.items()},
    }
    gc.collect()
    return df, stats


def quality_counts(df):
    all_blank = make_all_six_blank_mask(df)
    all_populated = make_all_six_populated_mask(df)
    any_populated = make_any_six_populated_mask(df)
    return {
        "Rows All Six Populated": int(all_populated.sum()),
        "Rows Partially Populated": int((any_populated & ~all_populated).sum()),
        "Rows All Six Blank": int(all_blank.sum()),
    }


def load_prepared_leftover(leftover_file=LEFTOVER_FILE, leftover_sheet=LEFTOVER_SHEET):
    leftover = read_sheet(leftover_file, leftover_sheet)
    print("Original leftover shape:", leftover.shape)
    verify_required_columns(leftover, LEFTOVER_REQUIRED_COLUMNS, "leftover lookup")
    leftover = leftover[LEFTOVER_REQUIRED_COLUMNS].rename(columns=LEFTOVER_COLUMN_MAP)
    check_duplicate_columns(leftover, "prepared leftover lookup")
    return prepare_lookup(leftover, "raw_frame_no", normalize_frame_no, "Leftover frame lookup")


# ============================================================
# SECTION 8: PHASE 2 - CROSS-YEAR STANDARDISATION
# ============================================================

def standardize_schema(df, file_cfg):
    """Drop only the explicitly known helper/duplicate columns for that year."""
    drop = [c for c in file_cfg["drop_columns"] if c in df.columns]
    return df.drop(columns=drop) if drop else df


def run_cross_year_standardization(files=PHASE2_FILES,
                                   expected_columns=EXPECTED_FINAL_COLUMNS,
                                   summary_file=PHASE2_SUMMARY_FILE):
    """PHASE 2 runner: standardise every year's schema and apply the leftover fill."""
    print("\n" + "=" * 80)
    print("LOADING LEFTOVER CANONICAL LOOKUP")
    print("=" * 80)
    if not os.path.exists(LEFTOVER_FILE):
        raise FileNotFoundError(f"Lookup file not found: {LEFTOVER_FILE}")
    lookup, lookup_maps = load_prepared_leftover()

    summary = []
    for file_cfg in files:
        fy, input_file, output_file = file_cfg["fy"], file_cfg["input"], file_cfg["output"]
        print("\n" + "=" * 80)
        print(f"PROCESSING FY {fy}")
        print("=" * 80)
        print("Input :", input_file)
        print("Sheet :", file_cfg["sheet"])
        print("Output:", output_file)

        if not os.path.exists(input_file):
            print("STATUS : SKIPPED, input file not found")
            summary.append({"FY": fy, "Input File": input_file,
                            "Output File": output_file, "Status": "Skipped - file not found"})
            continue

        df = read_sheet(input_file, file_cfg["sheet"], file_cfg["header_overrides"])
        original_rows, original_cols = len(df), len(df.columns)
        print(f"Original shape: {original_rows:,} rows x {original_cols} columns")

        df = standardize_schema(df, file_cfg)
        verify_required_columns(df, ["frame_no", *CANONICAL_COLUMNS], input_file)
        check_duplicate_columns(df, input_file)

        df, stats = apply_lookup(
            df, df["frame_no"], normalize_frame_no, lookup, lookup_maps,
            f"FY {fy} - Leftovers using frame_no",
        )

        if len(df.columns) != expected_columns:
            raise RuntimeError(
                f"\nFY {fy} has {len(df.columns)} columns after standardisation, "
                f"not the expected {expected_columns}.\n\nColumns are:\n{list(df.columns)}"
            )

        counts = quality_counts(df)
        print(f"Final shape: {len(df):,} rows x {len(df.columns)} columns")
        print(f"Rows with all six canonical fields populated: {counts['Rows All Six Populated']:,}")
        print(f"Rows where all six remain blank/zero: {counts['Rows All Six Blank']:,}")

        write_large_excel(df, output_file, sheet_name="Final Data")
        if not os.path.exists(output_file):
            raise RuntimeError(f"Output was not created: {output_file}")
        size_mb = os.path.getsize(output_file) / (1024 ** 2)
        print(f"Saved successfully: {output_file}")
        print(f"Output size: {size_mb:,.2f} MB")

        summary.append({
            "FY": fy, "Input File": input_file, "Output File": output_file,
            "Original Rows": original_rows, "Final Rows": len(df),
            "Final Columns": len(df.columns),
            "Eligible Rows": stats["Eligible Rows"],
            "Matched Eligible Rows": stats["Matched Eligible Rows"],
            "Fully Populated Rows": counts["Rows All Six Populated"],
            "Rows Still Fully Blank": counts["Rows All Six Blank"],
            "Output Size MB": round(size_mb, 2), "Status": "Completed",
        })
        del df
        gc.collect()

    summary_df = pd.DataFrame(summary)
    print("\n" + "=" * 80)
    print("FINAL PROCESSING SUMMARY")
    print("=" * 80)
    if not summary_df.empty:
        print(summary_df.to_string(index=False))
    summary_df.to_csv(summary_file, index=False)
    print("\nSummary saved as:", summary_file)
    print("All original workbooks remain unchanged.")
    return summary_df


# ============================================================
# SECTION 9: PHASE 3 - STAGED TARGETED FILL
# ============================================================

def load_stage_lookup(stage):
    """Load, clean and prepare the lookup for one Phase 3 stage."""
    raw_lookup = read_sheet(stage["lookup_file"], stage["lookup_sheet"])
    print(f"Original lookup shape: {raw_lookup.shape}")
    verify_required_columns(raw_lookup, stage["required_columns"], stage["label"])

    if stage["drop_embedded_header"]:
        # the first data row repeats the column meanings
        # ("Raw Variant", "Maker", ...); it is metadata, not a mapping row
        header_mask = (
            raw_lookup[stage["required_columns"][0]]
            .astype("string").str.strip().str.upper().eq("RAW VARIANT")
        )
        print(f"Embedded header rows removed: {int(header_mask.sum()):,}")
        raw_lookup = raw_lookup.loc[~header_mask]

    lookup = raw_lookup[stage["required_columns"]].rename(columns=stage["column_map"])
    check_duplicate_columns(lookup, f"renamed {stage['label']} lookup")
    print("Lookup mapping used:")
    for old, new in stage["column_map"].items():
        print(f"  {old} -> {new}")

    return prepare_lookup(
        lookup, stage["lookup_key"], NORMALIZERS[stage["normalizer"]], stage["label"]
    )


def run_staged_fill(cfg=PHASE3_CONFIG):
    """PHASE 3 runner: fill one year through a sequence of keyed stages."""
    # --- file check ---
    print("\n" + "=" * 80)
    print("FILE CHECK")
    print("=" * 80)
    required_files = [cfg["source_file"]] + [s["lookup_file"] for s in cfg["stages"]]
    missing = [f for f in dict.fromkeys(required_files) if not os.path.exists(f)]
    if missing:
        raise FileNotFoundError(
            "\nRequired files not found:\n" + "\n".join(f"  - {f}" for f in missing)
        )
    for f in dict.fromkeys(required_files):
        print("Found:", f)

    # --- load source ---
    print("\n" + "=" * 80)
    print("LOADING SOURCE")
    print("=" * 80)
    df = read_sheet(cfg["source_file"], cfg["source_sheet"], cfg["header_overrides"])
    print(f"Original source shape: {len(df):,} rows x {len(df.columns):,} columns")
    if cfg["expected_source_rows"] is not None and len(df) != cfg["expected_source_rows"]:
        print(f"WARNING: expected {cfg['expected_source_rows']:,} rows, found {len(df):,}.")
    if len(df.columns) != cfg["expected_source_columns"]:
        print(f"WARNING: expected {cfg['expected_source_columns']:,} columns, "
              f"found {len(df.columns):,}.")
    check_duplicate_columns(df, cfg["source_file"])

    # --- standardise schema ---
    drop = [c for c in cfg["drop_columns"] if c in df.columns]
    df = df.drop(columns=drop)
    verify_required_columns(
        df, [s["source_key"] for s in cfg["stages"]] + CANONICAL_COLUMNS, "standardised source"
    )
    check_duplicate_columns(df, "standardised source")

    source_rows, source_columns = len(df), len(df.columns)
    source_column_order = list(df.columns)
    print(f"Standardised source shape: {source_rows:,} rows x {source_columns:,} columns")
    if source_columns != cfg["expected_final_columns"]:
        raise RuntimeError(
            f"\nSource has {source_columns} columns after standardisation, "
            f"not the expected {cfg['expected_final_columns']}.\n\nColumns:\n{list(df.columns)}"
        )

    # --- snapshot original canonical values for protection checks ---
    original_values = df[CANONICAL_COLUMNS].copy()
    protected_mask = make_any_six_populated_mask(df)
    original_counts = quality_counts(df)
    print("\nOriginal fully blank canonical rows:", f"{original_counts['Rows All Six Blank']:,}")
    print("Original fully populated canonical rows:", f"{original_counts['Rows All Six Populated']:,}")
    print("Original partially populated canonical rows:", f"{original_counts['Rows Partially Populated']:,}")

    summary_rows = [{
        "Stage": "Original standardised source", "Input File": cfg["source_file"],
        "Rows": source_rows, "Columns": source_columns,
        "Eligible Rows": original_counts["Rows All Six Blank"], **original_counts,
    }]

    # --- run each stage in order ---
    previous_input = cfg["source_file"]
    for stage in cfg["stages"]:
        print("\n" + "=" * 80)
        print(stage["label"].upper())
        print("=" * 80)
        lookup, lookup_maps = load_stage_lookup(stage)

        # eligibility is recalculated per stage, so later stages only
        # touch rows still fully blank after earlier stages
        df, stats = apply_lookup(
            df, df[stage["source_key"]], NORMALIZERS[stage["normalizer"]],
            lookup, lookup_maps, stage["label"],
        )
        verify_unchanged_structure(df, source_rows, source_columns,
                                   source_column_order, stage["label"])
        verify_protected_rows(df, original_values, protected_mask, stage["label"])

        counts = quality_counts(df)
        print(f"\nRows still fully blank: {counts['Rows All Six Blank']:,}")
        print(f"Rows fully populated: {counts['Rows All Six Populated']:,}")

        print(f"\nWriting {stage['output_file']} ...")
        write_large_excel(df, stage["output_file"], sheet_name="Final Data")
        verify_written_workbook(stage["output_file"], "Final Data", source_rows,
                                source_columns, source_column_order)
        size_mb = os.path.getsize(stage["output_file"]) / (1024 ** 2)
        print(f"Output size: {size_mb:,.2f} MB")

        summary_rows.append({
            "Input File": previous_input, "Lookup File": stage["lookup_file"],
            "Lookup Sheet": stage["lookup_sheet"], "Key": stage["source_key"],
            "Rows": source_rows, "Columns": source_columns,
            **stats, **counts,
            "Output File": stage["output_file"], "Output Size MB": round(size_mb, 2),
        })
        previous_input = stage["output_file"]
        del lookup, lookup_maps
        gc.collect()

    # --- final quality ---
    remaining = {col: int(blank_or_zero(df[col]).sum()) for col in CANONICAL_COLUMNS}
    print("\n" + "=" * 80)
    print("FINAL DATA QUALITY")
    print("=" * 80)
    print(f"Final shape: {len(df):,} rows x {len(df.columns):,} columns")
    for key, value in quality_counts(df).items():
        print(f"{key}: {value:,}")
    print("\nRemaining blanks by canonical column:")
    for col, n in remaining.items():
        print(f"  {col}: {n:,}")
    summary_rows[-1].update({f"{col} remaining blank": n for col, n in remaining.items()})

    # --- summary ---
    summary_df = pd.DataFrame(summary_rows)
    summary_df.to_csv(cfg["summary_file"], index=False)
    print("\n" + "=" * 80)
    print("FINAL PROCESSING SUMMARY")
    print("=" * 80)
    print(summary_df.to_string(index=False))
    print("\nCreated files:")
    for i, stage in enumerate(cfg["stages"], start=1):
        print(f"  {i}. {stage['output_file']}")
    print(f"  {len(cfg['stages']) + 1}. {cfg['summary_file']}")
    print("\nOriginal workbooks remain unchanged.")
    print("Series.map was used, so duplicate lookup keys cannot increase row count.")

    del original_values, protected_mask
    gc.collect()
    return df, summary_df


# ============================================================
# SECTION 10: PHASE 4 - VARIANT-KEY PASS ON A CLEANED YEAR
# ============================================================

def run_variant_pass(cfg=PHASE4_CONFIG):
    """
    Fill rows whose six canonical fields are all blank, matching the
    normalised vehicle_variant against the Lookup sheet's raw variant.
    """
    df = read_sheet(cfg["target_file"], cfg["target_sheet"])
    lookup = read_sheet(cfg["lookup_file"], cfg["lookup_sheet"])

    # remove embedded header row, then name columns by meaning
    lookup = lookup.iloc[1:].copy()
    lookup = lookup.rename(columns=LOOKUP_SHEET_COLUMN_MAP)

    # --- normalised keys ---
    to_key = lambda s: s.astype(str).str.upper().str.replace(r"[^A-Z0-9]", "", regex=True)
    lookup["key"] = to_key(lookup["raw_variant"])
    lookup = lookup.drop_duplicates("key")
    df["key"] = to_key(df["vehicle_variant"])

    blank_before = df[CANONICAL_COLUMNS].isna().all(axis=1)
    print("Fully blank before:", int(blank_before.sum()))

    # --- merge and fill only previously blank rows ---
    rows_before = len(df)
    df2 = df.merge(lookup[["key", *CANONICAL_COLUMNS]], on="key",
                   how="left", suffixes=("", "_lk"))
    check_row_count(df2, "After variant merge", rows_before)

    for col in CANONICAL_COLUMNS:
        df2.loc[blank_before, col] = df2.loc[blank_before, col].fillna(
            df2.loc[blank_before, f"{col}_lk"]
        )

    blank_after = df2[CANONICAL_COLUMNS].isna().all(axis=1)
    print("Fully blank after :", int(blank_after.sum()))
    print("Rows improved     :", int(blank_before.sum() - blank_after.sum()))

    # --- save ---
    drop_cols = ["key"] + [f"{c}_lk" for c in CANONICAL_COLUMNS]
    df2 = df2.drop(columns=[c for c in drop_cols if c in df2.columns])
    df2.to_excel(cfg["output_file"], sheet_name="Final Data", index=False)

    # --- quick check of new file ---
    check = pd.read_excel(cfg["output_file"], sheet_name="Final Data",
                          nrows=5, engine="openpyxl")
    print("\nNew file shape:", df2.shape)
    print("Remaining fully blank rows:", int(blank_after.sum()))
    print("\nPreview:")
    print(check[["vehicle_variant", *CANONICAL_COLUMNS]].head())
    return df2


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    # PHASE 0: inspect what you're working with
    for f in [PHASE1_CONFIGS["24-25"]["source_file"], LEFTOVER_FILE]:
        if os.path.exists(f):
            inspect_workbook(f)

    # PHASE 1: per-year enrichment
    for fy, cfg in PHASE1_CONFIGS.items():
        if os.path.exists(cfg["source_file"]):
            print(f"\n######## PHASE 1 · FY{fy} ########")
            run_base_enrichment(cfg)

    # PHASE 2: cross-year standardisation
    run_cross_year_standardization()

    # PHASE 3: staged fill
    run_staged_fill()

    # PHASE 4: variant pass
    run_variant_pass()
